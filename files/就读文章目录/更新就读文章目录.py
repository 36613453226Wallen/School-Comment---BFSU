#!/usr/bin/env python3
"""读取文章目录.json并生成就读文章目录DOCX。"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "文章目录.json"
OUTPUT_PATH = ROOT / "就读文章目录.docx"


def set_font(run, name="宋体", size=9, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def compact(paragraph, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if align is not None:
        paragraph.alignment = align


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data.get("records") or []

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(data["title"])
    set_font(r, "黑体", 18, True, (31, 78, 121))

    pending = sum(1 for record in records if str(record.get("status", "")).startswith("待PDF"))
    p = doc.add_paragraph(
        f"登记 {len(records)} 篇｜待PDF {pending} 篇｜更新 {data['updated_at']}｜口令：“{data['trigger']}”"
    )
    compact(p, WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, "宋体", 8, color=(95, 99, 104))
    p.paragraph_format.space_after = Pt(8)

    headers = ["序号", "DOCX 全名", "时间", "学校", "专业", "状态"]
    widths = [0.5, 2.35, 1.45, 1.1, 1.1, 1.15]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "D9EAF7")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "黑体", 9, True, (31, 78, 121))

    if not records:
        row = table.add_row()
        row.cells[0].merge(row.cells[-1])
        cell = row.cells[0]
        cell.text = "暂无已收录文章。发送「收录就读文章」并附 PDF 与原文链接。"
        shade(cell, "FFF8E7")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "宋体", 9, color=(122, 88, 26))
    else:
        for row_index, record in enumerate(records, start=1):
            time_text = (
                f"发布：{record.get('published') or '未知'}\n"
                f"编辑：{record.get('edited') or '未知'}\n"
                f"快照：{record.get('snapshot') or '未知'}"
            )
            values = [
                str(record["id"]),
                record.get("docx") or "（尚未命名）",
                time_text,
                record.get("school") or "待辨认",
                record.get("major") or "待辨认",
                record.get("status") or "未知",
            ]
            row = table.add_row()
            for i, (value, width) in enumerate(zip(values, widths)):
                cell = row.cells[i]
                cell.width = Inches(width)
                cell.text = value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index % 2 == 0:
                    shade(cell, "F6F8FA")
                align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT
                compact(cell.paragraphs[0], align)
                for run in cell.paragraphs[0].runs:
                    set_font(run, "宋体", 8.5)
            if record.get("source_url"):
                source = row.cells[1].add_paragraph("原文：" + record["source_url"])
                compact(source)
                for run in source.runs:
                    set_font(run, "宋体", 7.5, color=(46, 116, 181))
            if record.get("analysis_docx"):
                extra = row.cells[1].add_paragraph("评注：" + record["analysis_docx"])
                compact(extra)
                for run in extra.runs:
                    set_font(run, "宋体", 7.5, color=(89, 89, 89))

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    compact(note)
    r = note.add_run(
        "维护规则：新增 PDF 后按加入顺序编号；首次发布、后续编辑与快照时间分开写。"
        "学校或专业只填文章里能确定的信息；侧面线索拿不准时先在对话里问，再改名。"
        "不得读取 Wushi/，除非另行授权。"
    )
    set_font(r, "宋体", 8, color=(89, 89, 89))

    doc.core_properties.title = data["title"]
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()
