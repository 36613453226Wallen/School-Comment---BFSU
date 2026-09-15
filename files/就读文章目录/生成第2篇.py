#!/usr/bin/env python3
"""Generate original and annotated DOCX for article 2 from the FireShot PDF."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "口令总表"))
from docx_style import (  # noqa: E402
    INK,
    NAVY,
    SLATE,
    add_bottom_border,
    compact,
    set_run_font,
)

ROOT = Path(__file__).resolve().parent
SOURCE_URL = "https://www.zhihu.com/question/308130958/answer/2048414048543306786"
TITLE = "北外就读啥体验？可能性抹杀、朝鲜语别报未凋零怪君"
ORIGINAL_NAME = f"2{TITLE}.docx"
ANNOTATED_NAME = f"2{TITLE}.分维评注.docx"

BLOCKS = [
    (
        "p",
        "本回答仅代表亚洲学院朝鲜语专业，也许其他的专业就读体验很好吧",
    ),
    (
        "p",
        "自从去年综评面试第一天到北外已经一年了，当年的我也通过知乎等平台不断的搜索这个学校怎么样，如今也终于能够写下自己的感受了。",
    ),
    ("h", "北外给我最大的感受在于对可能性的抹杀"),
    (
        "p",
        "北外只适合那些意向特别特别特别明确，且通过一定途径验证了自己意向的坚定性与天赋的人报考，除此以外不推荐任何人报考这所学校，如果不是因为一些专业相关的特殊原因，给我再来一次的机会我绝对不会选择这个学校。（本人当年放弃了西交，厦大来的你外，再来一次我哪怕去上外都不会来）永远永远不要高估自己对任何东西的热爱，永远不要高估他",
    ),
    (
        "p",
        "学校（清北除外）给你带来的title都是虚的，能够进入体制内的终究是少数中的少数，当你真正放到市场化就业的时候，你才会发现（除非面试官是这个学校的），北外和上外几乎没有差距，倒不如说上外由于地理优势好反而能有点优势，作为一个小语种学生，面试官眼里北外，上外，对外经贸等等是在同一档（也许一些强大对双非也在这档）不要妄想这个title会给你带来任何的优势。",
    ),
    (
        "p",
        "至于对于可能性的抹杀，则在于糟糕的课程安排和设计（注意是课程安排和设计，与授课老师无关，部分老师还是很好的，当然也有坏的）从入学之初便不断强调你们是xxx方向的，你们以后一定要去做xxx，所以现在就要做好准备（这个xxx是学术研究向的，我不明白一个专业将近40个人，最后能有几个人配做学术研究，大部分不还是市场化就业吗）安排的课程无用不说，部分课程十分古老（说的就是你Access+）与市场脱节十分严重就算了，老师还意识不到这一点，导致市场所需要的能力，你外培养了0。",
    ),
    (
        "p",
        "很难想象作为被ai影响最深的专业之一，居然在课程设计上完全与ai无关，甚至全校都没有ai相关的选修课，在专业课上老师要么对ai闭口不提，要么让我们用豆包改作业（这个豆包真的太难绷了）导致专业内绝大部分的学生对于ai的认知还停留在豆包这种对话式ai+。",
    ),
    (
        "p",
        "最难以理解的是都2026年，居然还有学校将实习视为洪水猛兽，目前的就业形式懂得都懂，没有实习在秋招是什么结果就不说了，而你外依旧保留暑假小学期与大二军训，同时大三甚至大四都一直有专业课的安排（我看到大四下还要上专业课我人都傻了）这就直接导致在学期内实习基本不可能，除非你能找到每周出勤两到三天的神仙公司（如果真的找到了记得@我，我也要去）而假期实习不仅要透支自己的假期，更要和其他二三线城市的好院校的同学竞争，那你北外在北京三环的地理优势的意义在哪呢，方便去买隔壁的绿色鸭腿吗？（翘课实习，请代课就别想了，小班教学，少一个都很明显，老师几乎认识专业内的所有学生，你跟老师求情也只会给你最官方的回答）",
    ),
    (
        "p",
        "最后聊聊你外引以为傲的人文关怀和官僚主义，给树织个毛衣，安个眼睛，摆点花花草草不叫人文关怀，真正为学生着想，打心眼里把学生当做独立的个体，而不是就业率或者其他率的数值，校外人员进校偷衣服的时候的人文关怀呢，被当成魏公村溜娃公园的时候的人文关怀呢，周末早上七八点彩排影响学生休息的时候的人文关怀呢，行政老师踢皮球，说什么都不让做，做不到，官僚主义尽显的人文关怀呢，让学生在10平米的宿舍后背贴后背，cosKappa，校园网30一个月卡得要死的时候的人文关怀呢，这还只是列举了一个学期的事情，若待满四年不知道会见证什么神奇生物（魏公村地铁站D口的某位没关系，我等你的老师点了个赞）宿舍的逆天舍友是一直都有的，不如说北外就是喜欢招这些逆天的人，换宿舍是不存在的，堵嘴是一定要有的，那我还能说什么呢，这篇文章又能存活多久呢。",
    ),
    (
        "p",
        "还有很多老生常谈的事情，如：食堂，优绩主义，精致利己，已经懒得再提了，想要了解的可以去别的帖子细看，就先写到这里吧，以后有机会会来补充的，最后提醒各位不要报考朝鲜语！宁愿选择提前批不接受调剂也不要选！",
    ),
    ("h", "2026.9.10更新"),
    (
        "p",
        "北外确实没有想象中的那么好，但是大部分专业课老师都是很好的，这点我在之前的回答也没有否认，希望这篇回答不要影响到很好的授课老师，我真的很喜欢很多老师",
    ),
]

POINTS = [
    {
        "name": "开篇限定范围",
        "claim": "只代表亚洲学院朝鲜语，并预留「别的专业也许很好」。",
        "support": "「本回答仅代表亚洲学院朝鲜语专业，也许其他的专业就读体验很好吧」。",
        "reason": "先把外推范围收窄，后面再骂「你外」，读的时候要记得这层括号。",
    },
    {
        "name": "可能性抹杀与再选后悔",
        "claim": "北外只适合意向极其明确、并验证过坚定性和天赋的人；否则不推荐。再选一次连上外都不来北外。不要高估热爱。",
        "support": "「给我再来一次的机会我绝对不会选择这个学校」；「放弃了西交，厦大来的你外」；「永远永远不要高估自己对任何东西的热爱」。",
        "reason": "这是全文总评。强度很高，但是第一人称后悔，不是录取数据。",
    },
    {
        "name": "title 在市场化就业里虚",
        "claim": "清北除外，学校 title 虚；体制内是少数；市场化就业里北外和上外几乎没差距，上外地理或更好；小语种里北外/上外/对外经贸同一档。",
        "support": "「学校（清北除外）给你带来的title都是虚的」；「北外和上外几乎没有差距」；「不要妄想这个title会给你带来任何的优势」。",
        "reason": "有就业场景（面试官、小语种档次），没有出处。「一些强大对双非也在这档」是压缩口语，不能当分类标准。",
    },
    {
        "name": "课程：学术导向、Access+、就业能力为零",
        "claim": "可能性被抹杀，是因为课程安排和设计糟糕，不是因为所有老师都差。入学就定学术 xxx 方向；近 40 人专业多数市场化就业；课无用且古老（点名 Access+）；市场要的能力北外培养了 0。",
        "support": "「注意是课程安排和设计，与授课老师无关」；「说的就是你Access+」；「你外培养了0」。",
        "reason": "这是最可核对的就读细节。Access+、学术口令、班额都具体。培养了 0 是情绪总评。",
    },
    {
        "name": "无 AI、豆包改作业",
        "claim": "朝鲜语被 AI 冲击大，课程完全无 AI；全校无 AI 选修；老师闭口或用豆包改作业；学生对 AI 停留在对话式豆包。",
        "support": "「被ai影响最深的专业之一」；「全校都没有ai相关的选修课」；「用豆包改作业」。",
        "reason": "2026 年写的现场吐槽，信息新。是否「全校没有」需要别的材料核对，这里只当作者见闻。",
    },
    {
        "name": "实习被挡：小学期、军训、大四专业课",
        "claim": "学校仍把实习当洪水猛兽；暑假小学期+大二军训；大四下还有专业课；学期内实习几乎不可能；假期还要和二三线好学校抢；三环优势被讽成买隔壁绿色鸭腿。小班逃课/代课行不通。",
        "support": "「大四下还要上专业课我人都傻了」；「方便去买隔壁的绿色鸭腿吗？」；「翘课实习，请代课就别想了」。",
        "reason": "课表结构、实习冲突写得很具体。绿色鸭腿是魏公村玩笑，不是地理结论。",
    },
    {
        "name": "人文关怀 vs 官僚、宿舍与堵嘴",
        "claim": "给树织毛衣、安眼睛、摆花不是关怀。点名偷衣服、魏公村溜娃、周末早彩排、行政踢皮球、10㎡宿舍后背贴后背 cosKappa、校园网 30 一个月。换宿舍不存在，堵嘴一定有。",
        "support": "一串「……的人文关怀呢」；「堵嘴是一定要有的，那我还能说什么呢，这篇文章又能存活多久呢」。",
        "reason": "生活槽点可核对，也带火药。「招这些逆天的人」是泄愤，不当生源结论。地铁 D 口那句像在点某个账号，转写保留，不解释成事实。",
    },
    {
        "name": "不要报考朝鲜语，以及给老师留余地",
        "claim": "结尾劝退朝鲜语，宁选提前批也不接受调剂。9 月 10 日更新强调：学校没那么好，但大部分专业课老师很好，希望别伤到老师。",
        "support": "「最后提醒各位不要报考朝鲜语！」；「希望这篇回答不要影响到很好的授课老师」。",
        "reason": "劝退是立场，不是录取建议书。更新段把「课」和「老师」拆开，和前文「与授课老师无关」一致，这点中肯。",
    },
]


def add_blank_lines(doc, n=4):
    for _ in range(n):
        p = doc.add_paragraph("＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿")
        compact(p, after=2, line=1.4)
        for run in p.runs:
            set_run_font(run, "宋体", 11, color=RGBColor(0xBB, 0xBB, 0xBB))


def new_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    return doc


def add_kicker(doc, text):
    p = doc.add_paragraph()
    compact(p, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, "黑体", 10, True, NAVY)


def add_title(doc, text):
    p = doc.add_paragraph()
    compact(p, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_bottom_border(p)
    run = p.add_run(text)
    set_run_font(run, "黑体", 16, True, NAVY)


def add_meta(doc):
    lines = [
        "作者：未凋零怪君",
        "知乎问题：北京外国语大学就读是怎样体验？",
        "编辑于：2026-09-10 19:43 · 北京（截图页脚）",
        "首次发布：未知（截图只显示编辑时间，不补造）",
        "PDF 快照：2026-09-15 11:29 UTC（FireShot 元数据）",
        "171 人赞同（界面数字，已从正文剔除）",
    ]
    for line in lines:
        p = doc.add_paragraph()
        compact(p, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(line)
        set_run_font(run, "宋体", 9, color=SLATE)


def add_body(doc):
    for kind, text in BLOCKS:
        p = doc.add_paragraph()
        if kind == "h":
            compact(p, before=10, after=8, line=1.2)
            add_bottom_border(p, "8A4B12", "8")
            run = p.add_run(text)
            set_run_font(run, "黑体", 13, True, NAVY)
        else:
            compact(p, after=8, line=1.35, first_line=22)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            run = p.add_run(text)
            set_run_font(run, "宋体", 12, color=INK)


def add_link(doc):
    p = doc.add_paragraph()
    compact(p, before=12, after=0, line=1.15)
    add_bottom_border(p, "D6DEE8", "8")
    run = p.add_run("知乎原文：")
    set_run_font(run, "黑体", 10, True, NAVY)
    run = p.add_run(SOURCE_URL)
    set_run_font(run, "宋体", 10, color=RGBColor(0x2E, 0x74, 0xB5))


def add_analysis(doc):
    h = doc.add_paragraph()
    compact(h, before=16, after=8, line=1.1)
    add_bottom_border(h, "8A4B12", "12")
    run = h.add_run("分维评注（副本加写，原文未改）")
    set_run_font(run, "黑体", 14, True, NAVY)

    intro = doc.add_paragraph()
    compact(intro, after=8, line=1.2)
    run = intro.add_run(
        "作者大约从 8 个方面评价北外，但反复声明只代表亚洲学院朝鲜语。"
        "下面每条先写分论点，再写支撑句和理由，最后单独看中肯程度。"
        "这是整理，不是替作者改口。原文里的「你外」、小写 ai、Access+ 都按截图保留。"
    )
    set_run_font(run, "宋体", 11, color=INK)

    for i, point in enumerate(POINTS, start=1):
        p = doc.add_paragraph()
        compact(p, before=8, after=4, line=1.15)
        run = p.add_run(f"{i}. {point['name']}")
        set_run_font(run, "黑体", 12, True, NAVY)
        for label, key in (("分论点", "claim"), ("支撑", "support"), ("理由", "reason")):
            q = doc.add_paragraph()
            compact(q, after=3, line=1.2)
            run = q.add_run(f"{label}：")
            set_run_font(run, "黑体", 11, True, NAVY)
            run = q.add_run(point[key])
            set_run_font(run, "宋体", 11, color=INK)

    judge = doc.add_paragraph()
    compact(judge, before=10, after=6, line=1.2)
    run = judge.add_run("分论点是否中肯")
    set_run_font(run, "黑体", 13, True, NAVY)

    verdict = (
        "整体中肯程度：作为「朝鲜语一年级综评生」的怒贴，中等偏上；作为全校总评，外推过猛。"
        "开篇限定、课程/实习/小学期/军训/小班、宿舍网费、更新段给老师留余地，都有可核对的细节。"
        "title 虚、北外上外同一档、培养了 0、全校没有 AI 选修，是观感，目录里不当统计事实。"
        "「不要报考朝鲜语」是劝退，不是专业评估报告。"
        "「你外」是作者对北外的讽刺叫法，转写不改成「北外」。"
        "截图是整页图片，个别短语（如「一些强大对双非也在这档」、地铁 D 口点赞那句）按所见字形保留。"
    )
    p = doc.add_paragraph()
    compact(p, after=10, line=1.3)
    run = p.add_run(verdict)
    set_run_font(run, "宋体", 11, color=INK)

    blank_title = doc.add_paragraph()
    compact(blank_title, before=8, after=4, line=1.1)
    run = blank_title.add_run("讨论留白（给读者自己写）")
    set_run_font(run, "黑体", 13, True, NAVY)

    prompts = [
        "如果只看朝鲜语，哪几条最像课表事实？哪几条只是心情？",
        "「与授课老师无关」和结尾「很喜欢很多老师」，能不能同时成立？",
        "实习、AI、title 这三件，你想反驳、补充，还是换一个专业对照？",
    ]
    for prompt in prompts:
        p = doc.add_paragraph()
        compact(p, after=2, line=1.15)
        run = p.add_run(prompt)
        set_run_font(run, "楷体", 11, color=SLATE)
        add_blank_lines(doc, 3)


def write_original():
    doc = new_doc()
    add_kicker(doc, "就读文章  ·  原文转写")
    add_title(doc, TITLE)
    add_meta(doc)
    add_body(doc)
    add_link(doc)
    path = ROOT / ORIGINAL_NAME
    doc.core_properties.title = TITLE
    doc.save(path)
    print(f"已生成：{path}")
    return path


def write_annotated():
    doc = new_doc()
    add_kicker(doc, "就读文章  ·  分维评注副本")
    add_title(doc, TITLE)
    add_meta(doc)
    add_body(doc)
    add_analysis(doc)
    add_link(doc)
    path = ROOT / ANNOTATED_NAME
    doc.core_properties.title = TITLE + "（分维评注）"
    doc.save(path)
    print(f"已生成：{path}")
    return path


if __name__ == "__main__":
    write_original()
    write_annotated()
