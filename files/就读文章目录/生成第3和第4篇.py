#!/usr/bin/env python3
"""Generate original and annotated DOCX for articles 3 and 4 (北外 MJC)."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "口令总表"))
from docx_style import (  # noqa: E402
    INK,
    NAVY,
    SLATE,
    add_bottom_border,
    compact,
    set_run_font,
)

ROOT = Path(__file__).resolve().parent


def add_blank_lines(doc, n=4):
    for _ in range(n):
        p = doc.add_paragraph("＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿")
        compact(p, after=2, line=1.4)
        for run in p.runs:
            set_run_font(run, "宋体", 11, color=RGBColor(0xBB, 0xBB, 0xBB))


def new_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    return doc


def add_kicker(doc, text):
    p = doc.add_paragraph()
    compact(p, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, "黑体", 10, True, NAVY)


def add_title(doc, text):
    p = doc.add_paragraph()
    compact(p, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_bottom_border(p)
    run = p.add_run(text)
    set_run_font(run, "黑体", 16, True, NAVY)


def add_lines(doc, lines, *, size=9, color=SLATE, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=None, font="宋体"):
    for line in lines:
        p = doc.add_paragraph()
        compact(p, after=after, line=1.15, align=align, first_line=first_line)
        if first_line:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.35
        run = p.add_run(line)
        set_run_font(run, font, size, color=color)


def add_link(doc, text):
    p = doc.add_paragraph()
    compact(p, before=12, after=0, line=1.15)
    add_bottom_border(p, "D6DEE8", "8")
    run = p.add_run("知乎原文：")
    set_run_font(run, "黑体", 10, True, NAVY)
    run = p.add_run(text)
    set_run_font(run, "宋体", 10, color=RGBColor(0x2E, 0x74, 0xB5))


def add_analysis(doc, intro, points, verdict, prompts):
    h = doc.add_paragraph()
    compact(h, before=16, after=8, line=1.1)
    add_bottom_border(h, "8A4B12", "12")
    run = h.add_run("分维评注（副本加写，原文未改）")
    set_run_font(run, "黑体", 14, True, NAVY)

    p = doc.add_paragraph()
    compact(p, after=8, line=1.2)
    run = p.add_run(intro)
    set_run_font(run, "宋体", 11, color=INK)

    for i, point in enumerate(points, start=1):
        q = doc.add_paragraph()
        compact(q, before=8, after=4, line=1.15)
        run = q.add_run(f"{i}. {point['name']}")
        set_run_font(run, "黑体", 12, True, NAVY)
        for label, key in (("分论点", "claim"), ("支撑", "support"), ("理由", "reason")):
            r = doc.add_paragraph()
            compact(r, after=3, line=1.2)
            run = r.add_run(f"{label}：")
            set_run_font(run, "黑体", 11, True, NAVY)
            run = r.add_run(point[key])
            set_run_font(run, "宋体", 11, color=INK)

    judge = doc.add_paragraph()
    compact(judge, before=10, after=6, line=1.2)
    run = judge.add_run("分论点是否中肯")
    set_run_font(run, "黑体", 13, True, NAVY)
    p = doc.add_paragraph()
    compact(p, after=10, line=1.3)
    run = p.add_run(verdict)
    set_run_font(run, "宋体", 11, color=INK)

    blank_title = doc.add_paragraph()
    compact(blank_title, before=8, after=4, line=1.1)
    run = blank_title.add_run("讨论留白（给读者自己写）")
    set_run_font(run, "黑体", 13, True, NAVY)
    for prompt in prompts:
        p = doc.add_paragraph()
        compact(p, after=2, line=1.15)
        run = p.add_run(prompt)
        set_run_font(run, "楷体", 11, color=SLATE)
        add_blank_lines(doc, 3)


def write_pair(spec):
    title = spec["title"]
    original = ROOT / f"{spec['id']}{title}.docx"
    annotated = ROOT / f"{spec['id']}{title}.分维评注.docx"

    doc = new_doc()
    add_kicker(doc, "就读文章  ·  原文转写")
    add_title(doc, title)
    add_lines(doc, spec["meta"])
    add_lines(doc, spec["body"], size=12, color=INK, after=8, align=None, first_line=22)
    add_link(doc, spec["link"])
    doc.core_properties.title = title
    doc.save(original)
    print(f"已生成：{original}")

    doc = new_doc()
    add_kicker(doc, "就读文章  ·  分维评注副本")
    add_title(doc, title)
    add_lines(doc, spec["meta"])
    add_lines(doc, spec["body"], size=12, color=INK, after=8, align=None, first_line=22)
    add_analysis(doc, spec["intro"], spec["points"], spec["verdict"], spec["prompts"])
    add_link(doc, spec["link"])
    doc.core_properties.title = title + "（分维评注）"
    doc.save(annotated)
    print(f"已生成：{annotated}")


ARTICLE3 = {
    "id": 3,
    "title": "北外中政MJC咋选？国新出名叫北外、走211选中政冰洁ice",
    "link": "未随截图提供。问题：北京外国语大学和中国政法大学的 MJC 哪个好一些？",
    "meta": [
        "作者：冰洁ice",
        "知乎问题：北京外国语大学和中国政法大学的 MJC 哪个好一些？",
        "首次发布：2019-05-04 09:46（截图「发布于」）",
        "编辑时间：未知（截图未显示编辑）",
        "PDF 快照：用户 2026-09-15 补入的回答截图，无 FireShot 元数据",
        "3 人赞同（界面数字，已从正文剔除）",
    ],
    "body": [
        "北外的话只有国新比较出名，其他不好说",
        "中国政法大学今年是第二年开设mjc",
        "考察难度应该不高 且报考人数还没有那么多",
        "和学硕比要简单一些",
        "你想在这两个学校的mjc中选择，建议：",
        "如果英语很好 想往国际发展 选北外",
        "如果只想走个211而已 且英语不是强项 选中政",
    ],
    "intro": "作者大约从 4 个方面比较北外和中政的 MJC。下面每条先写分论点，再写支撑句和理由。这是 2019 年的短评，不是 2026 年的招生简章。",
    "points": [
        {
            "name": "北外只有国新出名",
            "claim": "北外这边，比较拿得出手的是国际新闻（国新），别的不好说。",
            "support": "「北外的话只有国新比较出名，其他不好说」。",
            "reason": "一句话印象，没有课程或就业数据。目录里只当 2019 年口碑，不当学院排名。",
        },
        {
            "name": "中政 MJC 当时还新、人少、比学硕简单",
            "claim": "截图那年是中政开 MJC 的第二年；考察不难、报考的人还不多；比学硕简单。",
            "support": "「今年是第二年开设mjc」；「考察难度应该不高 且报考人数还没有那么多」；「和学硕比要简单一些」。",
            "reason": "「今年」锚定在发布日 2019-05-04，不能读成现在。难度和人数是作者判断。",
        },
        {
            "name": "英语好、想做国际传播 → 北外",
            "claim": "英语强、目标往国际发展，选北外。",
            "support": "「如果英语很好 想往国际发展 选北外」。",
            "reason": "和北外国新、语言优势的常见说法一致，但是没有举例。",
        },
        {
            "name": "只要 211、英语不强 → 中政",
            "claim": "如果只是要一个 211，而且英语不是强项，选中政。",
            "support": "「如果只想走个211而已 且英语不是强项 选中政」。",
            "reason": "把中政写成「走 211」通道，口气很直。两校当时都是 211，这条是策略建议，不是办学评价。",
        },
    ],
    "verdict": (
        "整体中肯程度：作为 2019 年的择校短信，中等。分流标准清楚（英语/国际 vs 只要 211），但证据几乎全是印象。"
        "「国新出名」「中政第二年」「人少好考」都没有出处，而且已过多年。"
        "正文很短，没有就读日常，更像报考建议而不是校园体验。"
        "3 个赞、2 条评论是界面数字，不说明建议被广泛验证。"
    ),
    "prompts": [
        "2019 年的「中政第二年、人少」，到现在还成立吗？",
        "「只要走个 211」这个标准，你接受还是反感？",
        "英语和国际传播之外，还有哪条是这篇没写到的？",
    ],
}

ARTICLE4 = {
    "id": 4,
    "title": "上岸北外MJC导师咋选？越早联系越不被动7.3倍的C6H12O6",
    "link": "未随截图提供。问题：本人今年上岸北外mjc，想问关于导师的选择是开学前自己提前联系，还是等到开学之后再选？",
    "meta": [
        "作者：7.3倍的C6H12O6　签名：科研狗",
        "知乎问题：本人今年上岸北外mjc，想问关于导师的选择是开学前自己提前联系，还是等到开学之后再选？",
        "问题补充：想毕业后直接工作的话，有没有师哥师姐推荐的老师（提问者的话，不是回答正文）",
        "首次发布：2020-05-29 10:35（截图「发布于」）",
        "编辑时间：未知（截图未显示编辑）",
        "PDF 快照：2026-09-15 22:20 UTC（FireShot 元数据）",
        "谢邀 @梅子酱（界面邀请，已从正文结构保留一句）",
    ],
    "body": [
        "越早联系越好，原因如下：",
        "（1）提前占大牛组的坑。按规定每个老师每年最多指导3名硕士生，早联系，可以提前让老师了解你，并把你收入麾下。如果等到开学，你的理想导师收够人，会很遗憾！",
        "（2）可以提前对导师进行了解。有充裕的时间阅读课题组的文献，并对以后从事方向有大概了解！如果不喜欢导师所有的研究方向，换课题组也不是不可以（双选制）！",
        "（3）开学后联系会很被动！开学后，会有很多事情要忙，如果提前联系好导师，可以马上入组找座位等，甚至暑假就可以提前入学！",
        "查阅导师资料、了解导师为人、选择对的导师，这些比考研重要！",
    ],
    "intro": "作者用 3 条理由主张开学前就联系导师，最后把选导师抬到比考研还重要。同页还有 DearAliciaaa 的短答，不并入这篇原文。",
    "points": [
        {
            "name": "总判断：越早越好",
            "claim": "导师要尽早联系，不要等到开学再选。",
            "support": "开篇「越早联系越好，原因如下」。",
            "reason": "先给结论再列原因，结构清楚。没有写自己是不是北外 MJC 在读。",
        },
        {
            "name": "名额：每老师每年最多 3 个硕士",
            "claim": "早联系是为了占大牛组名额；开学再去，心仪导师可能已经收满。",
            "support": "「按规定每个老师每年最多指导3名硕士生」；「理想导师收够人，会很遗憾」。",
            "reason": "「按规定」说得很硬，但截图没有附学院文件。名额因年、因导师而异，目录里不当现行规章。",
        },
        {
            "name": "了解方向，不合适可以双选换组",
            "claim": "提前读课题组文献，摸清方向；不喜欢也可以换组，因为是双选。",
            "support": "「换课题组也不是不可以（双选制）」。",
            "reason": "把「双选」写进建议，比只抢大牛完整。换组实际难度没有写。",
        },
        {
            "name": "开学后再联系会被动，暑假甚至可提前入学",
            "claim": "开学后事情多、很被动；提前联系好可以马上入组找座位，暑假就能提前入学。",
            "support": "「开学后联系会很被动」；「甚至暑假就可以提前入学」。",
            "reason": "入组、占座、暑假入学是具体场景，因导师而异。提问者还问「毕业直接工作推荐哪位老师」，这篇没有答。",
        },
        {
            "name": "选对导师比考研重要",
            "claim": "查资料、看人品、选对人，比考研本身更重要。",
            "support": "结尾「这些比考研重要！」。",
            "reason": "是态度，不是比较数据。对已经上岸的人有用，对还在备考的人容易吓到。",
        },
    ],
    "verdict": (
        "整体中肯程度：作为「要不要提前联系」的经验帖，中等偏上；作为北外 MJC 官方流程说明，不够。"
        "三条理由自洽：名额、了解、开学后被动。"
        "「每年最多 3 名硕士」「暑假提前入学」需要对照当年学院规定，不能直接当成 2026 年规则。"
        "作者没回答「想就业该跟谁」；同页 DearAliciaaa 写「有的老师接受提前联系，有的不一定」，正好是限制条件，读合 2 时两篇对着看更完整。"
        "广告条已从正文剔除。"
    ),
    "prompts": [
        "你所在学院现在还有没有「每导师每年 3 个硕士」这种硬名额？",
        "想就业和想读博，提前联系的话术该不该一样？",
        "DearAliciaaa 说「有的老师不一定接受提前联系」，会不会把这篇的「越早越好」打薄？",
    ],
}


if __name__ == "__main__":
    write_pair(ARTICLE3)
    write_pair(ARTICLE4)
