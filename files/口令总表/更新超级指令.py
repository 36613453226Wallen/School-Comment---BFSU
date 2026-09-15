#!/usr/bin/env python3
"""Generate the formatted *超级指令.docx catalog."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from docx_style import (
    INK,
    NAVY,
    SLATE,
    add_bottom_border,
    compact,
    shade_cell,
    set_cell_margins,
    set_run_font,
)


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "超级指令.json"
OUTPUT_PATH = ROOT / "*超级指令.docx"


def add_text(paragraph, text, *, name="宋体", size=10.5, bold=False, color=INK):
    run = paragraph.add_run(text)
    set_run_font(run, name, size, bold, color)
    return run


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    commands = data["commands"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    kicker = doc.add_paragraph()
    compact(kicker, after=2, line=1.0)
    add_text(kicker, "School-Comment---BFSU  ·  口令总表", name="黑体", size=10, bold=True, color=NAVY)

    title = doc.add_paragraph()
    compact(title, after=4, line=1.1)
    add_bottom_border(title, "1F4E79", "16")
    add_text(title, data["title"], name="黑体", size=20, bold=True, color=NAVY)

    intro = doc.add_paragraph()
    compact(intro, before=6, after=8, line=1.15)
    add_text(
        intro,
        "本文件汇总本仓库已经生成过的关键词口令、提示词文档和自动更新规则。"
        "新指令不以聊天宣布为准，必须以本表新增或更新的条目为准。",
        size=10.5,
        color=INK,
    )

    rule = doc.add_paragraph()
    compact(rule, after=10, line=1.15)
    add_text(rule, "自动更新：", name="黑体", size=10.5, bold=True, color=NAVY)
    add_text(rule, data["auto_update_rule"], size=10.5, color=INK)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["序号", "口令 / 提示词", "类型", "用途"]
    widths = [0.6, 2.4, 1.2, 2.7]
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = ""
        shade_cell(cell, "1F4E79")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        compact(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, header, name="黑体", size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for index, command in enumerate(commands, start=1):
        row = table.add_row()
        values = [
            str(command["id"]),
            command["keyword"],
            command["type"],
            command["purpose"],
        ]
        for i, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.text = ""
            shade_cell(cell, "F4F7FA" if index % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            compact(p, after=0, line=1.08, align=align)
            add_text(p, value, name="宋体", size=9, color=INK)

    spacer = doc.add_paragraph()
    compact(spacer, after=4, line=1.0)

    for command in commands:
        heading = doc.add_paragraph()
        compact(heading, before=10, after=4, line=1.1)
        add_bottom_border(heading, "1F4E79", "10")
        add_text(heading, f"{command['id']}. {command['keyword']}", name="黑体", size=13, bold=True, color=NAVY)

        meta = doc.add_paragraph()
        compact(meta, after=4, line=1.1)
        add_text(meta, f"{command['type']}  ·  建立于 {command['created']}", size=9, color=SLATE)

        fields = [
            ("用途", command["purpose"]),
            ("来源", command["source"]),
            ("输入", command["inputs"]),
            ("操作", command["operation"]),
            ("输出", command["outputs"]),
            ("备注", command.get("notes") or "无"),
        ]
        for label, value in fields:
            p = doc.add_paragraph()
            compact(p, after=4, line=1.18)
            add_text(p, f"{label}：", name="黑体", size=10.5, bold=True, color=NAVY)
            add_text(p, value, size=10.5, color=INK)

    footer = doc.add_paragraph()
    compact(footer, before=12, after=0, line=1.15)
    add_text(
        footer,
        f"共 {len(commands)} 条｜更新 {data['updated_at']}｜数据源 超级指令.json｜生成脚本 更新超级指令.py",
        size=9,
        color=SLATE,
    )

    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "口令与提示词总表"
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()
