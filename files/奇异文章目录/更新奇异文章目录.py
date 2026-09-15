#!/usr/bin/env python3
"""读取奇异目录.json并生成奇异文章目录DOCX。"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "奇异目录.json"
OUTPUT_PATH = ROOT / "奇异文章目录.docx"


def set_font(run, name="宋体", size=9, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def compact(paragraph, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if align is not None:
        paragraph.alignment = align


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data.get("records") or []

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data["title"])
    set_font(r, "黑体", 18, True, (122, 48, 108))

    p = doc.add_paragraph(
        f"登记 {len(records)} 篇｜更新 {data['updated_at']}｜口令：“{data['trigger']}”｜与就读文章目录序号独立"
    )
    compact(p, WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, "宋体", 8, color=(95, 99, 104))
    p.paragraph_format.space_after = Pt(8)

    headers = ["序号", "DOCX 全名", "时间", "状态"]
    widths = [0.6, 3.4, 2.0, 1.4]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "F3E5F5")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "黑体", 9, True, (122, 48, 108))

    if not records:
        row = table.add_row()
        row.cells[0].merge(row.cells[-1])
        cell = row.cells[0]
        cell.text = "暂无奇异文章。发送「收录奇异文章」并附 PDF 与原文链接。"
        shade(cell, "FFF8E7")
        compact(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            set_font(run, "宋体", 9, color=(122, 88, 26))
    else:
        for row_index, record in enumerate(records, start=1):
            time_text = (
                f"发布：{record.get('published') or '未知'}\n"
                f"编辑：{record.get('edited') or '未知'}"
            )
            values = [
                str(record["id"]),
                record.get("docx") or "（尚未命名）",
                time_text,
                record.get("status") or "未知",
            ]
            row = table.add_row()
            for i, (value, width) in enumerate(zip(values, widths)):
                cell = row.cells[i]
                cell.width = Inches(width)
                cell.text = value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index % 2 == 0:
                    shade(cell, "FBF4FC")
                align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
                compact(cell.paragraphs[0], align)
                for run in cell.paragraphs[0].runs:
                    set_font(run, "宋体", 8.5)
            if record.get("source_url"):
                source = row.cells[1].add_paragraph("原文：" + record["source_url"])
                compact(source)
                for run in source.runs:
                    set_font(run, "宋体", 7.5, color=(46, 116, 181))

    doc.core_properties.title = data["title"]
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()
