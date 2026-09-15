#!/usr/bin/env python3
"""Generate Prompt-就读文章半自动收录流程.docx with original, blank, suggested, and gray source."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x21, 0x25, 0x29)
NAVY = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x88, 0x88, 0x88)
SLATE = RGBColor(0x5F, 0x63, 0x68)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "Prompt-就读文章半自动收录流程.docx"
SOURCE_COPY = Path(__file__).resolve().parent / "收录就读文章-原版提示词.md"

ORIGINAL = """PS文件外操作.记录prompt到Files/记录

口令1：收录就读文章

目标：把我手动保存的文章（这次是学校学生体验相关）PDF转换为可编辑DOCX，并同步维护文章目录。

每次执行时：
1. 接收新PDF和原文链接；将PDF保存到 files/ZHIHU ，不得读取 WuShi/，除非我另行授权。
2. 优先使用PDF内容。识别正文、一级/二级标题、引用、编号列表和发布时间；去掉网页导航、私信数、广告及评论区。
3. 生成接近原文排版的DOCX。文件名只保留“序号＋正文标题（标题包含简短问题和总评价）”，例如：
   - 1北外就读体验，后悔但又有点.docx
   - 2理工就读体验，还行中规中矩.docx【别按我的格式，来点趣味和花样】
4.
4.1在docx分析原文，原文docx和原文docx标题不变。
4.2生成相同内容的副本，并在副本docx文件上创作：指出作者有几个分论点（有几个方面评价学校，或者叫维度。），以及支持分论点的句子和理由。目的是给读者以清晰化。最后给出分论点的评价是否中肯。并留下讨论的留白区域供用户或者读者进行分析和记录。
4.3在DOCX结尾写入对应的知乎原文链接。
5. 分别记录首次发布时间、后续编辑时间和PDF快照时间；截图只显示到年或月时不得补造日期。
6. 更新 文章目录相关的json（如没有请生成），并重新生成姚勇文章目录.docx。目录保持简洁，至少显示序号、DOCX全名、时间、学校名字（如文章里有）+相关专业（如果文章里有涉及，有侧面信息但是无法判断的可以在cursor界面提出疑问，让用户进行更名）和状态。
7. 若PDF尚未出现或无法读取，不访问受限网页补造正文；保留序号和链接，状态写“待PDF”，并明确缺少什么。

口令2：“修改口令”
 调口令1的提示词给我，并等待我回复
"""

BLANK = """口令1：【短触发词】

目标：把【材料类型】PDF 转成可编辑 DOCX，并同步维护【目录名】。

每次执行时：
1. 接收新 PDF 和原文链接；保存到【归档目录】，不得读取【禁读目录】，除非另行授权。
2. 优先使用 PDF。识别正文、标题、引用、列表和发布时间；去掉导航、广告和评论区。
3. 生成接近原文排版的 DOCX。文件名：【命名规则】。标题要有趣味，不套固定句式。
4.1 原文 DOCX 与标题保持不动。
4.2 另存分维评注副本：分论点 / 支撑句 / 是否中肯 / 讨论留白。
4.3 DOCX 结尾写原文链接。
5. 分别记录首次发布、后续编辑、快照时间；信息不全不补造。
6. 更新【json】并重生成【目录docx】。目录字段：【字段列表】。拿不准的学校/专业先提问。
7. 缺 PDF 时状态写「待PDF」，不访问受限网页补正文。

口令2：修改口令
调出口令1完整提示词，等待回复后再改。

并列口令：【收录奇异文章】——独立序号、独立目录。
"""

SUGGESTED = """口令1：收录就读文章

目标：把手动保存的学校学生体验 PDF 转成可编辑 DOCX，并同步维护《就读文章目录》。

每次执行时：
1. 接收新 PDF 和原文链接；保存到 files/ZHIHU/。不得读取 Wushi/（WuShi/），除非另行授权。
2. 优先使用 PDF。识别正文、一级/二级标题、引用、编号列表和发布时间；去掉网页导航、私信数、广告及评论区。
3. 生成接近原文排版的原文 DOCX。文件名只保留「序号＋正文标题」。标题须含短问句和总评价，并且有趣味，例如：
   - 1北外这趟，悔过但还想再咬一口.docx
   - 2理工日子，不惊不喜刚刚好.docx
4.1 原文 DOCX 与原文标题保持不动。
4.2 另存内容相同的分维评注副本（建议：序号标题.分维评注.docx）：标出有几个评价维度、支撑句和理由、是否中肯，并留讨论空白。
4.3 两份 DOCX 结尾都写知乎原文链接。
5. 分别记录首次发布时间、后续编辑时间和 PDF 快照时间；截图只到年或月时不补日期。
6. 更新 files/就读文章目录/文章目录.json，运行 更新就读文章目录.py，重写《就读文章目录.docx》（不是姚勇文章目录）。目录字段：序号、DOCX 全名、时间、学校、专业、状态。学校或专业无法判断时在对话里提问，不擅自改名。
7. PDF 缺失或读不出时，不访问受限网页补正文；保留序号和链接，状态写「待PDF」，并写明缺什么。

口令2：修改口令
原样调出「收录就读文章」提示词，然后等待回复。确认前不改口令。

口令3：收录奇异文章
对应女娲库「收录姚勇奇闻」。奇闻/吐槽/脑洞材料放入 files/奇异文章目录/，独立序号，更新 奇异目录.json。

口令4：更新目录
半自动刷新 files/目录-就读部署/ 对话 Menu。平时不自动改文件。
"""


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, size=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "黑体", size, True, NAVY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "1F4E79")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_body(doc, text, color=INK, size=11):
    for block in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(block if block else " ")
        set_run_font(run, "宋体", size, False, color)


def generate():
    SOURCE_COPY.write_text(ORIGINAL, encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("School-Comment---BFSU")
    set_run_font(run, "黑体", 10, True, NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Prompt-就读文章半自动收录流程")
    set_run_font(run, "黑体", 20, True, NAVY)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("含原版、挖空版、建议版及深灰色原文  ·  对应女娲库根目录 Prompt 文档")
    set_run_font(run, "宋体", 10, False, SLATE)

    add_heading(doc, "原版")
    add_body(doc, ORIGINAL)

    add_heading(doc, "挖空版")
    add_body(doc, BLANK)

    add_heading(doc, "建议版")
    add_body(doc, SUGGESTED)

    add_heading(doc, "深灰色原文")
    add_body(doc, ORIGINAL, color=GRAY, size=10)

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(16)
    run = footer.add_run(
        "母本：提示词记录/提示词记录.docx。"
        "口令总表：files/口令总表/*超级指令.docx。"
        "并列口令「收录奇异文章」已写入建议版与超级指令。"
    )
    set_run_font(run, "宋体", 9, False, SLATE)

    doc.core_properties.title = "Prompt-就读文章半自动收录流程"
    doc.save(OUTPUT)
    print(f"已生成：{OUTPUT}")


if __name__ == "__main__":
    generate()
