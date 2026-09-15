#!/usr/bin/env python3
"""从目录数据.json生成紧凑的DOCX对话目录。"""

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "目录数据.json"
OUTPUT_PATH = ROOT / "目录-就读评论之Menu.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_paragraph(paragraph, *, size=8.5, bold=False, align=None):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(size)
        run.bold = bold


def generate():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data["records"]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(data["title"])
    run.font.name = "黑体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(15)
    run.bold = True

    note = doc.add_paragraph(
        f"共 {len(records)} 条｜更新：{data['updated_at']}｜策略：{data['update_mode']}"
    )
    format_paragraph(note, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    note.paragraph_format.space_after = Pt(4)

    headers = ["标号", "问题总结", "答案概要（大概完成了什么）", "使用模型", "对话时间"]
    widths = [0.45, 2.15, 4.55, 1.6, 1.55]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    header_row = table.rows[0]
    header_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, (header, width) in enumerate(zip(headers, widths)):
        cell = header_row.cells[idx]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_width(cell, width)
        set_cell_shading(cell, "D9EAF7")
        format_paragraph(
            cell.paragraphs[0],
            size=8.5,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, record in enumerate(records, start=1):
        row = table.add_row()
        values = [
            str(record["id"]),
            record["question"],
            record["answer"],
            record["model"],
            record.get("time") or "约/未知",
        ]
        for idx, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[idx]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, width)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F6F8FA")
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT
            format_paragraph(cell.paragraphs[0], size=8.2, align=align)

    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "就读评论仓库对话目录"
    doc.save(OUTPUT_PATH)
    print(f"已生成：{OUTPUT_PATH}")


if __name__ == "__main__":
    generate()
