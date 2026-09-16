#!/usr/bin/env python3
"""Build 论坛风云: first-person weight folder, UG/PG splits, remaining-count Excel.

Does not rewrite 论坛纷纭 sample texts. Copies them into
「学院名本科部 / 学院名研究生部」subfolders (max 7 each).
自述权重+ 目标每部 7–10 条第一人称（如有才放）；高翻本科不适用。
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FENYUN = Path("/workspace/files/论坛纷纭")
ROOT = Path("/workspace/files/论坛风云")
WEIGHT = ROOT / "自述权重+"
RETRIEVED = "2026-09-16"
BAR = 5  # 业余印象分析门槛：每格合格第一人称就读条数
PEOPLE_BAR = 3
FOLDER_MIN = 7  # 用户要求每部 7–10 条；如有才放，不编造
FOLDER_MAX = 10

COLLEGES = ["英语学院", "俄语学院", "高级翻译学院", "国际商学院", "国际关系学院"]

# 自述权重+：第一人称对该院该层次。
# 合格 = 独立论坛/书评就读现场，非约稿。
# 半合格 = 身份对得上，但考研/复试/课表转载/入学后两三句，或约稿里密课业第一人称。
# 半合格进夹凑 7–10 条，不充合格门槛。高翻本科不适用，不凑。
WEIGHT_SAMPLES = [
    {
        "college": "英语学院",
        "dept": "本科部",
        "weight": "合格",
        "title": "豆瓣-英院幸存者偏差",
        "handle": "匿名豆瓣楼主",
        "people": "豆瓣英院楼主",
        "src": "论坛纷纭/英语学院/03-本科-豆瓣小组.txt",
        "url": "https://www.douban.com/group/topic/269635697/",
        "note": "在读第一人称，分班考、保送生、2016课改、精读词数。",
    },
    {
        "college": "英语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "金榜转载-英专大二课表",
        "handle": "自称大二，未署名",
        "people": "金榜英专大二",
        "src": "论坛纷纭/英语学院/04-本科-金榜教育转载问答.txt",
        "url": "https://m.360jibing.com/news/409412.html",
        "note": "课表细，但转载站、大三以后是观摩。只算半合格。",
    },
    {
        "college": "俄语学院",
        "dept": "本科部",
        "weight": "合格",
        "title": "就读文章5-Whever俄语后悔",
        "handle": "Whever",
        "people": "Whever",
        "src": "论坛纷纭/俄语学院/01-本科-知乎.txt",
        "url": "https://www.zhihu.com/question/64391382/answer/3331909814",
        "note": "仓库第5篇。短，但是俄语本科第一人称。",
    },
    {
        "college": "俄语学院",
        "dept": "本科部",
        "weight": "合格",
        "title": "就读文章6-Whever劝退俄语",
        "handle": "Whever",
        "people": "Whever",
        "src": "论坛纷纭/俄语学院/02-本科-知乎.txt",
        "url": "https://www.zhihu.com/question/597450083/answer/3000551110",
        "note": "同一人第二篇。推荐国关，不是国关学生自述。",
    },
    {
        "college": "俄语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "考研论坛-2018俄语MTI复试现场",
        "handle": "旧帖楼主（站点改版后用户名被导航淹没）",
        "people": "2018俄语MTI复试楼主",
        "src": "论坛纷纭/俄语学院/04-研究生部-考研论坛-kaoyan-com.txt",
        "url": "http://bbs.kaoyan.com/t8648738p1",
        "note": "考场第一人称，不是入学后课堂日记。",
    },
    {
        "college": "国际商学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "经管之家-曹渊国商大四考金专",
        "handle": "caoyuanibsbfsu / 曹渊",
        "people": "曹渊",
        "src": "论坛纷纭/国际商学院/01-本科-研究生部-经管之家-人大经济论坛.txt",
        "url": "https://bbs.pinggu.org/thread-4577562-1-1.html",
        "note": "身份是国商大四，正文几乎全是考研科目。",
    },
    {
        "college": "国际商学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "经管之家-lu-llaby66北外金融本科推免",
        "handle": "lu-llaby66",
        "people": "lu-llaby66",
        "src": "论坛纷纭/国际商学院/02-本科-经管之家.txt",
        "url": "https://bbs.pinggu.org/thread-4945124-1-1.html",
        "note": "「我本科就是在北外读金融」，有实习，仍偏保研流程。",
    },
    {
        "college": "国际商学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "查字典转载-金专入学后午餐会宿舍CFA",
        "handle": "未署名，总分400",
        "people": "金专400分帖",
        "src": "论坛纷纭/国际商学院/04-研究生部-查字典公务员网转载考研经验.txt",
        "url": "https://gwy.chazidian.com/news385566/",
        "note": "「来到北外之后」几句生活，转载站可能美化。",
    },
    {
        "college": "国际商学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "文书帮转载-国商学习两个多月",
        "handle": "未署名（与保研帖高度相似，或同一批转载）",
        "people": "文书帮国商两个月",
        "src": "",
        "url": "https://www.wenshubang.com/xindetihui/2895526.html",
        "quote": (
            "总之，来到了梦想中的北外还是挺高兴的。北外地理位置相当不错，在西三环上，旁边的院校有北理、中青政、民大等，"
            "离人大、北大、清华也很近，交通非常便利。北外校园小巧精致，硬件设施一直在改善，图书馆和国际大厦修建的很高大上。"
            "同时，北外学风浓厚，思想包容开放，社团活动相当丰富……"
            "在国商学习了两个多月后，我深刻地感受到老师很认真负责，为学生着想。"
            "学院的学术导师和外界聘请的企业导师都很牛，对今后的学习和工作都有很大的帮助。"
            "总之，来到北外是个不悔的选择。"
        ),
        "note": "续搜新增。入学后短评，转载站，可能与经管之家保研帖同源。",
    },
    {
        "college": "国际关系学院",
        "dept": "本科部",
        "weight": "合格",
        "title": "转载-外交系在读双专业",
        "handle": "未署名",
        "people": "njarts外交系在读",
        "src": "论坛纷纭/国际关系学院/02-本科-教育问答转载.txt",
        "url": "https://www.njarts.cn/a_jiaoyu/202106/216899.html",
        "note": "「我是北外外交系在讀的」。短，转载站。",
    },
    {
        "college": "国际关系学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "考研帮-wjb0911外交学毕业回看",
        "handle": "wjb0911",
        "people": "wjb0911",
        "src": "论坛纷纭/国际关系学院/05-研究生部-考研帮.txt",
        "url": "https://yz.kaoyan.com/bfsu/jingyan/566fb7c0d6c56.html",
        "note": "自称国关研究生国家奖学金两年。抒情多、课业少。",
    },
    {
        "college": "英语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "北鼎-英美文学网面90.6",
        "handle": "未署名，2020英美文学",
        "people": "2020英美文学90.6",
        "src": "论坛纷纭/英语学院/11-研究生部-北鼎教育.txt",
        "url": "https://www.beidingedu.cn/sys-nd/8697.html",
        "note": "复试第一人称，机构页。不是入学后课堂。",
    },
    {
        "college": "俄语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "听力室转载-2016俄语口译复试",
        "handle": "未署名，自称2016俄语口译研究生",
        "people": "2016俄语口译复试楼主",
        "src": "论坛纷纭/俄语学院/12-研究生部-听力室-学习网转载.txt",
        "url": "http://ru.tingroom.com/xuexi/eyxiff/19641.html",
        "note": "考场视译流程具体。转载站。不是入学后课堂。",
    },
    {
        "college": "英语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "Insight-康惠阳peer-pressure",
        "handle": "康惠阳",
        "people": "康惠阳",
        "src": "论坛纷纭/英语学院/01-本科-Insight-BFSU.txt",
        "url": "http://share.office369.com/yingyong/356932.html",
        "note": "约稿。精读、presentation、peer pressure 是课业现场，但不进合格。",
    },
    {
        "college": "英语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "Insight-吴旨瑨原著讨论",
        "handle": "吴旨瑨",
        "people": "吴旨瑨",
        "src": "论坛纷纭/英语学院/02-本科-Insight-BFSU.txt",
        "url": "http://share.office369.com/yingyong/356932.html",
        "note": "约稿。原著、理论文本、小组汇报。学长口吻。",
    },
    {
        "college": "英语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "校友会征文-赵培85级精读听力",
        "handle": "赵培",
        "people": "赵培",
        "src": "",
        "url": "https://xyh.bfsu.edu.cn/info/1371/17032.htm",
        "quote": (
            "印象最深的是上课听不懂。由于当时高中不重视听力和口语，所以学生的听说能力普遍薄弱，难以适应全英文的授课，"
            "记得当时教听力课的是一位白白净净的小帅哥，名字好像叫建华，上课时他总让大家听完录音后回答问题，我感到压力山大，听不懂就根本无法回答老师提出的问题。"
            "精读课老师习惯上课时在教室里不停地走动，随时叫起一位同学回答问题。课文难度很大，生词量很多，老师的讲解又听不懂，真是提心吊胆，生怕老师提问我。"
            "听力尚且如此，口语就更不必说了，班里同学的水平也是参差不齐，有一些同学来自北外附中和杭州外语学校，他们的口语很好，我的Speaking pair就是来自杭州外校的……"
            "经过了大概一个学期的痛苦适应期，我的听说能力有了突飞猛进的提高。"
            "到了大二，我们按照专业倾向分班，我分到了翻译班。"
        ),
        "note": "英语系1985级征文，课堂第一人称密，但是毕业三十多年后回看。",
    },
    {
        "college": "英语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "新祥旭-文学方向课全都上过",
        "handle": "未署名学长学姐口吻",
        "people": "新祥旭文学方向",
        "src": "论坛纷纭/英语学院/07-研究生部-新祥旭考研.txt",
        "url": "http://www.xxxedu.net/articlenew.php?id=321690",
        "note": "机构稿。「英美文学史、英美文学选读」像课表，无法核学籍。",
    },
    {
        "college": "英语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "查字典-美国研究复试拟录取",
        "handle": "未署名；西南211理科本科",
        "people": "美国研究三跨",
        "src": "论坛纷纭/英语学院/14-研究生部-查字典考研网.txt",
        "url": "https://ky.chazidian.com/kaoyan/544404/",
        "note": "复试题具体。原页521，不是入学后课堂。",
    },
    {
        "college": "英语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "路灯转载-跨考英语言文学复试",
        "handle": "未署名；自称本科不是学语言",
        "people": "路灯跨考英语言文学",
        "src": "论坛纷纭/英语学院/12-研究生部-路灯在职研究生转载.txt",
        "url": "https://www.125yan.com/zixun/25212.html",
        "note": "在职研站点洗稿。复试问答，不是读研日记。",
    },
    {
        "college": "英语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "北鼎-工科跨考611",
        "handle": "自称工科跨考生",
        "people": "北鼎工科跨考",
        "src": "论坛纷纭/英语学院/06-研究生部-北鼎教育.txt",
        "url": "https://www.beidingedu.cn/sys-nd/14424.html",
        "note": "机构学员稿。写的是考进来之前。",
    },
    {
        "college": "俄语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "招生办-陆瑶一周俄语",
        "handle": "陆瑶",
        "people": "陆瑶",
        "src": "",
        "url": "https://m.027art.com/beijingbenke/HTML/4785213.html",
        "quote": (
            "19级 俄语 陆瑶\n"
            "从高中“跳槽”到大学，我感受到的是自由，我们成了时间的主人，更应合理规划，在仅有一次的四年时光里做更多有意义的事情。"
            "缘分这东西真是妙不可言，和北外相识是在那张小小的校报，似乎是感受到了北外的召唤，我抓住了它，终有幸成为俄语学院的一名学生。"
            "这些时日身边的人似乎都在说俄语难，而经历一周学习后的我想说：一切的难易都是相对。"
            "“塞翁失马，焉知非福”，也许正是这种意义上的难，让我能够磨砺意志、提高学习能力，塑造一个更好的自己！"
            "东西院穿梭的日子让我见证了北外并非印象中的小……最后有个小期望：地下通道赶紧修好哇，这样同学们东西穿梭就更安全啦。"
        ),
        "note": "招生办新生采访。原页今日404，引文来自检索快照。入学一周，课业几乎没有。",
    },
    {
        "college": "俄语学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "豆瓣书评-08级东方俄语小白鼠",
        "handle": "S. Vågslid",
        "people": "S. Vågslid",
        "src": "",
        "url": "https://book.douban.com/subject/34998170/",
        "quote": (
            "作为第一批见到这本书的北外人，我们是被实验的小白鼠。。\n"
            "现在看到的这一册除了封面和我们08级用的一样，内容上改动好大。。"
            "大家都喊着说这本书难，你们不知道其实它的原版更难。。"
            "我们是唯一见过并且保存着这全国少数人才有的绝版。。。我不知道该高兴还是囧。。。"
        ),
        "note": "独立书评，点名08级和教材难度。太短，只算半合格。",
    },
    {
        "college": "俄语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "本科院校风采-刘岚普俄语笔译课",
        "handle": "刘岚普",
        "people": "刘岚普",
        "src": "",
        "url": "https://wyxy.hntou.edu.cn/zsjy_1305/xyfc/202004/t20200406_30413.html",
        "quote": (
            "大家好，我是13级俄语专业毕业生刘岚普，现在北外读翻译硕士，笔译方向。"
            "对一些人来说，研究生的第一年可能会很不适应。研一上学期，课程的安排无非就是一个接着一个的翻译课。"
            "从一个话题翻译到另外一个话题，感觉所有的课程都没有什么用。"
            "翻来覆去的就是翻译一些没有什么难度的东西，课堂上老师也只是就一些翻译难点讲解一下，有时甚至会觉得老师的译文还没有自己的好。"
            "一学期下来可能也讲不了几篇文章，自己的翻译能力似乎也没有什么提升。平时的作业量也不是很多，更多的时间都是交给自己处理。"
            "到了研一下学期，课程更加密集，各种针对性的训练也更多。"
            "我们专业课的老师经常说，翻译无非就是汉译汉，如果连汉语的意思都理解的不正确，那怎么能把它正确的翻译成俄语。"
            "在这里，很多人为了准备翻译资格证考试都是早出晚归。沉迷图书馆，不想学习，无所事事真的会有负罪感。"
        ),
        "note": "本科院校学长稿。课业现场密，但是约稿。汉译俄，记入俄语学院而非高翻。",
    },
    {
        "college": "俄语学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "校友会-赵梦雪硕博口译练习",
        "handle": "赵梦雪",
        "people": "赵梦雪",
        "src": "论坛纷纭/俄语学院/14-研究生部-北外校友会.txt",
        "url": "https://xyh.bfsu.edu.cn/info/1311/9432.htm",
        "note": "约稿出色个例。本科黑龙江大学。北外段有上课、论文、口译练习、苗澍最新材料。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "官网访谈-张旭复语班否决项",
        "handle": "张旭",
        "people": "张旭",
        "src": "论坛纷纭/高级翻译学院/10-研究生部-高翻学院官网-校友访谈.txt",
        "url": "https://gsti.bfsu.edu.cn/info/1511/3791.htm",
        "note": "约稿。作业否决项和《感觉身体被掏空》改编词是目前最密的高翻课业现场。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "校友会-文苑李长栓200字reference",
        "handle": "文苑",
        "people": "文苑",
        "src": "",
        "url": "https://xyh.bfsu.edu.cn/info/1311/9692.htm",
        "quote": (
            "在高翻学院读研的两年，第一年交传，第二年同传。她回忆说，"
            "“进入高翻学院的学生都是精英，学院的教学是高起点的，刚开始上课没有一点过渡，直接就是口译，大家一下子难以适应。"
            "尤其是听力，对没有口译经验的学生来说难度太大。"
            "那时候高翻学院的学生是各院系研究生中唯一每天自觉自习到晚上十点的群体，由于高强度的学习，气氛紧张压抑，大家没有时间参加社团活动。"
            "文苑的努力更甚旁人，每天学习14到16小时，一个月后，学习开始变得得心应手。"
            "她举了当时的笔译老师李长栓教授的例子，当时他布置作业的强度最大，一星期布置的作业只有200字的翻译，但是每一处翻译都要给出证据（reference），学生往往要花一周的时间完成。"
        ),
        "note": "校友会记者稿，出色个例。200字reference是课业细节。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "官网专访-尚姝辰研一口笔译",
        "handle": "尚姝辰",
        "people": "尚姝辰",
        "src": "论坛纷纭/高级翻译学院/05-研究生部-高翻学院官网-毕业生专访.txt",
        "url": "https://gsti.bfsu.edu.cn/info/1491/17501.htm",
        "note": "约稿。研一「高强度口笔译训练」一句，正文主要是求职。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "北鼎-研一依旧有演讲课",
        "handle": "未署名学姐口吻",
        "people": "北鼎高翻演讲课",
        "src": "",
        "url": "https://www.beidingedu.cn/sys-nd/14047.html",
        "quote": (
            "在各个课堂上老师都有强调，其实我们做翻译，不仅要能对语言的理解跟表达有深刻的认识，"
            "还要能在各种场合自如地表达自己的观点，所以我们研一依旧会有演讲课，这说明了对表达能力的重视。"
            "举个例子，我们的口译老师也在课堂上提到过，可能大家都知道气候变化。"
            "但是如果在复试的过程中提问气候变化背后的原因机制到底是什么，却很少有同学能够讲清楚。"
        ),
        "note": "机构页。研一演讲课、口译老师课堂各一句。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "考研论坛-19高翻MA一战",
        "handle": "旧帖楼主（站点改版后用户名被导航淹没）",
        "people": "19高翻MA一战",
        "src": "论坛纷纭/高级翻译学院/06-研究生部-考研论坛-kaoyan-com.txt",
        "url": "http://bbs.kaoyan.com/t10009599p1",
        "note": "备考帖。分数和题型具体，入学后课堂几乎没有。",
    },
    {
        "college": "高级翻译学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "李长栓文-匿名期末总结笔译助口译",
        "handle": "匿名高翻学生期末总结",
        "people": "高翻期末总结匿名",
        "src": "",
        "url": "https://www.jamoxi.com/thread/20190917/11288842.html",
        "quote": (
            "学生在校期间写的期末总结也经常提到这一点比如：\n"
            "李老师在课上反复强调“笔译是口译的基础”。我赞同这一观点"
            "笔译练习一则能够帮助我们扩充背景知识，有效减轻口译时的压力"
            "二则能够提升我们的语言质量，使得我们产出的译文能够更加简洁、凝练"
            "对于这两点，我深有体会每次口译课上的练习涉及到一些专业性较强的知识时，"
            "笔译作业过程中查到的背景知识都为我理解语篇提供了很大的帮助"
            "此外，对翻译作业的反复修改使我学会了如何用更加简洁的方式组织语言、搭建句子大大缩短了交传所需时间。"
        ),
        "note": "教授论文转引匿名期末总结。课业第一人称，但不是独立论坛日记。转载站OCR有错字。",
    },
    {
        "college": "国际商学院",
        "dept": "本科部",
        "weight": "合格",
        "title": "豆瓣-国商会计财管培养方案",
        "handle": "🍊🍊🍊",
        "people": "豆瓣国商会计楼主",
        "src": "",
        "url": "https://www.douban.com/group/topic/179507086/",
        "quote": (
            "楼主来自北外商学院，受学校背景限制，会计/财管的培养方案与财经类院校可能差距较大（比如注重英语轻视专业教学）……\n"
            "就北外而言，会计/财管对数学要求不高，学习内容基本参照考研数三的范围，不过也有一些文科同学认为学起来有困难。"
            "除此以外需要学习基础的统计学，会带一些实操的内容，如excel/stata等。"
            "英语课基本可以根据名字猜课程内容，包括听力课（听力材料包括BBC/VOA/TED/艾伦秀等）/口语课（纠音与对话/演讲/辩论）/"
            "写作课（基础写作/分析型/应用型）/综合课（类似中学），听力口语写作课课程质量很高，多为英文授课，对于口语较弱的同学会比较吃力。"
            "受学校背景影响，专业课程多为英文教材/课件+中/英文授课……"
            "会计信息系统：会计信息系统AIS和企业资源规划系统ERP，理论与实操，勉强能够上手金蝶/用友，贴近实务。"
            "整体上看，楼主认为有一定理解难度（不如工科），也有一定的记忆难度（不如文科），找工作或许需要一些“sense”。"
        ),
        "note": "独立豆瓣科普帖。课表级细节强，志愿向，不是后悔日记。",
    },
    {
        "college": "国际商学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "考研帮-生姜饼干金专复试",
        "handle": "生姜饼干",
        "people": "生姜饼干",
        "src": "论坛纷纭/国际商学院/03-研究生部-考研帮.txt",
        "url": "https://yz.kaoyan.com/bfsu/jingyan/597ab45236fe7.html",
        "note": "二外考入国商。复试形式逐年变，课堂日常少。",
    },
    {
        "college": "国际商学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "经管之家-simple_kaka国际商务拟录取",
        "handle": "simple_kaka",
        "people": "simple_kaka",
        "src": "论坛纷纭/国际商学院/13-研究生部-经管之家.txt",
        "url": "https://bbs.pinggu.org/thread-4945117-1-1.html",
        "note": "拟录取经验。入学后只一句免修英语传闻。",
    },
    {
        "college": "国际关系学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "中青报-2021级外交学德语法语",
        "handle": "未在本仓库核到真名",
        "people": "中青报2021外交学",
        "src": "论坛纷纭/国际关系学院/01-本科-中国青年报客户端-中青在线.txt",
        "url": "http://news.cyol.com/gb/articles/2024-06/03/content_v6wm5QIlg7.html",
        "note": "党媒约稿。英语文献、德法语课、实习比赛都有，出色个例。",
    },
    {
        "college": "国际关系学院",
        "dept": "本科部",
        "weight": "半合格",
        "title": "灯塔留学-Anette外交学法语",
        "handle": "Anette",
        "people": "Anette",
        "src": "",
        "url": "https://www.91goodschool.com/news/13143-10330998.html",
        "quote": (
            "最后被调剂到北外的外交学（法语与法语国家研究）专业，它实际上是主修外交学辅修法语的双学位专业，刚入学就要学法语，而且和法语专业的学生是同样的要求。"
            "进入大学后，面对完全陌生的法语学习，作为一个内向的人，我很难在陌生人面前开口说蹩脚的法语，害怕出丑。"
            "加上学院老师提醒我们是这个专业的最后一届学生，挂科就没有学位证，压力很大。那段时间又是疫情封校……"
            "因为记不住单词，又不敢开口说法语，成绩一直很差，法语总是不及格……"
            "大三上学期，我重新探索了学习方法，逐渐掌握了学习技巧，努力把大一大二落下的知识补起来，成绩有了明显提升，也恢复了对法语学习与国际关系研究的探索热情。"
            "本科期间因为学业太忙，尤其是法语课程和毕业论文占用了大量时间，没有办法同时准备雅思和申请材料。"
        ),
        "note": "留学顾问约稿。挂科、封校、法语课是就读现场，后半是申请广告。",
    },
    {
        "college": "国际关系学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "Free考研-wjb0911外交学复试",
        "handle": "wjb0911",
        "people": "wjb0911",
        "src": "论坛纷纭/国际关系学院/12-研究生部-Free考研转考研论坛.txt",
        "url": "http://school.freekaoyan.com/bj/bfsu/jingyan/20131104/1383568783197909.shtml",
        "note": "与毕业回看同一人。更早、更接近考场。独立叙述者不增加。",
    },
    {
        "college": "国际关系学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "刀豆文库-外交学复试跨专业",
        "handle": "未署名",
        "people": "刀豆外交学复试",
        "src": "论坛纷纭/国际关系学院/13-研究生部-刀豆文库转载.txt",
        "url": "https://www.daodoc.com/fanwen/qitafanwen/2204892.html",
        "note": "像外交学国际经济方向。原页打不开库，只留检索快照。",
    },
    {
        "college": "国际关系学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "北鼎-国关无参考书真题体会",
        "handle": "未署名；自称一家之言",
        "people": "北鼎国关真题",
        "src": "论坛纷纭/国际关系学院/06-研究生部-北鼎教育.txt",
        "url": "https://www.beidingedu.cn/sys-nd/11693.html",
        "note": "机构页。备考体会，不是导师组日常。",
    },
    {
        "college": "国际关系学院",
        "dept": "研究生部",
        "weight": "半合格",
        "title": "北鼎-外交学国际经济网面",
        "handle": "自称成功学员",
        "people": "北鼎外交学国际经济网面",
        "src": "论坛纷纭/国际关系学院/08-研究生部-北鼎教育.txt",
        "url": "https://www.beidingedu.cn/sys-nd/8742.html",
        "note": "机构学员稿。网面题具体。国际经济方向不要和国商金融搞混。",
    },
]

EXCLUDED_PROMO = [
    ("英语学院", "本科部", "Insight 康惠阳、吴旨瑨", "已进半合格。约稿出色个例，不充合格。"),
    ("英语学院", "本科部", "今日头条就业实录", "媒体成稿，未进夹。"),
    ("英语学院", "本科部", "英学职沙龙刘闯/林启和", "学院报道转述，不是主讲人第一人称原文。"),
    ("英语学院", "研究生部", "北鼎/新祥旭/路灯/查字典上岸稿", "已进半合格。考研机构稿，不充合格。"),
    ("高级翻译学院", "研究生部", "张旭、文苑、尚姝辰", "已进半合格。官网/校友会风采，不充合格。"),
    ("俄语学院", "研究生部", "赵梦雪校友会采访", "已进半合格。本科黑龙江大学；出色个例。"),
    ("俄语学院", "本科部", "豆瓣小语种辩护", "学校层，未写俄语学院。"),
    ("俄语学院", "本科部", "出国培训部俄语预科张可", "预科≠俄语学院统招。"),
    ("俄语学院", "研究生部", "新祥旭就业率100%", "机构推销。"),
    ("高级翻译学院", "本科部", "推免方案、保研指南", "无普通本科就读对象。"),
    ("高级翻译学院", "研究生部", "李长栓教师访谈", "老师看学生，不是学生自述；学生期末总结另条半合格。"),
    ("高级翻译学院", "研究生部", "北鼎法语口译在读邵炜", "法语学院，不记入高翻。"),
    ("国际商学院", "本科部", "iMAP郑宽、国际本科评价、冠城EAP李想", "留学/国际项目≠统招。"),
    ("国际商学院", "本科部", "高顿倪子君双学位申哥大", "ACCA机构学霸稿。"),
    ("国际商学院", "研究生部", "新祥旭MIB、Reddit申请帖", "备考/申请，不是就读。"),
    ("国际关系学院", "本科部", "中青报2021、Anette灯塔", "已进半合格。党媒/留学顾问约稿，不充合格。"),
    ("国际关系学院", "本科部", "Whever国关养老", "俄语生外推，作者不在国关。"),
    ("国际关系学院", "本科部", "豆瓣亚运村小樱花外交科普", "某校国关，评论像定福庄/外院，不能记北外国关。"),
    ("国际关系学院", "研究生部", "北鼎复试/真题", "已进半合格。机构备考。"),
    ("校级未分院", "—", "就读文章1 MouradBey、7 成就你的梦想", "未写五院之一；第7篇是付费咨询"),
    ("他院", "—", "就读文章2 朝鲜语、3–4 MJC", "亚洲学院/国新，不占五院格子"),
]


def classify_fenyun(name: str) -> str | None:
    stem = name[:-4] if name.endswith(".txt") else name
    if name == "本夹说明.txt" or stem.startswith("本夹"):
        return None
    if "研究生部" in stem and "本科" not in stem:
        return "研究生部"
    if "本科" in stem:
        return "本科部"
    if "研究生" in stem:
        return "研究生部"
    return None


def copy_fenyun_split():
    """Place existing 论坛纷纭 samples into 学院名本科部 / 学院名研究生部, max 7."""
    summary = {}
    for college in COLLEGES:
        src_dir = FENYUN / college
        buckets = {"本科部": [], "研究生部": []}
        for p in sorted(src_dir.glob("*.txt")):
            if p.name in ("本夹说明.txt",):
                continue
            dept = classify_fenyun(p.name)
            if dept:
                buckets[dept].append(p)
        summary[college] = {}
        for dept, files in buckets.items():
            dest = src_dir / f"{college}{dept}"
            dest.mkdir(parents=True, exist_ok=True)
            taken = files[:7]
            for old in dest.glob("*.txt"):
                if old.name != "本夹说明.txt":
                    old.unlink()
            for f in taken:
                shutil.copy2(f, dest / f.name)
            leftover = max(0, 7 - len(taken))
            (dest / "本夹说明.txt").write_text(
                f"{college}{dept}\n"
                f"从「论坛纷纭/{college}」原取样再放置，不删原文件。\n"
                f"本夹 {len(taken)} 条（上限 7，如有则放满；现有不足不编造）。\n"
                f"距每部 7 条还差 {leftover} 条公开取样。\n"
                f"其中约稿、机构稿、备考帖仍在，不代表自述权重+。\n"
                f"合格第一人称就读见 files/论坛风云/自述权重+/\n",
                encoding="utf-8",
            )
            summary[college][dept] = {"placed": len(taken), "short_of_7": leftover, "files": [f.name for f in taken]}
    return summary


def write_weight_files():
    for college in COLLEGES:
        for dept in ("本科部", "研究生部"):
            folder = WEIGHT / college / f"{college}{dept}"
            folder.mkdir(parents=True, exist_ok=True)
            for old in folder.glob("*.txt"):
                old.unlink()

    by_cell: dict[tuple[str, str], list] = {}
    for sample in WEIGHT_SAMPLES:
        key = (sample["college"], sample["dept"])
        by_cell.setdefault(key, []).append(sample)

    for college in COLLEGES:
        for dept in ("本科部", "研究生部"):
            folder = WEIGHT / college / f"{college}{dept}"
            items = by_cell.get((college, dept), [])[:FOLDER_MAX]
            for i, sample in enumerate(items, start=1):
                if sample.get("src"):
                    src = Path("/workspace/files") / sample["src"]
                    text = src.read_text(encoding="utf-8") if src.exists() else ""
                else:
                    text = (
                        f"学院：{sample['college']}\n"
                        f"层次：{dept}\n"
                        f"用户/署名：{sample['handle']}\n"
                        f"链接：{sample['url']}\n"
                        f"检索：{RETRIEVED} 续搜\n\n"
                        f"摘录：\n{sample.get('quote', '')}\n"
                    )
                header = (
                    f"自述权重+　{sample['weight']}\n"
                    f"标题：{sample['title']}\n"
                    f"独立叙述标记：{sample['people']}\n"
                    f"说明：{sample['note']}\n"
                    f"{'=' * 40}\n"
                )
                (folder / f"{i:02d}-{sample['weight']}-{sample['title']}.txt").write_text(
                    header + text, encoding="utf-8"
                )
            n = len(items)
            extra = ""
            na = college == "高级翻译学院" and dept == "本科部"
            if na:
                extra = "高翻公开建制以研究生为主，本科格按「不适用」计，不凑 7–10 条假自述。\n"
            short = 0 if na else max(0, FOLDER_MIN - n)
            target_line = (
                "本夹不适用普通本科就读对象，不设 7–10 条目标。\n"
                if na
                else (
                    f"合格/半合格第一人称：{n} 条（目标 {FOLDER_MIN}–{FOLDER_MAX}，如有才放，不编造）。\n"
                    f"距下限 {FOLDER_MIN} 还差 {short} 条公开取样。\n"
                )
            )
            (folder / "本夹说明.txt").write_text(
                f"{college}{dept}　自述权重+\n"
                f"{target_line}"
                f"约稿密课业、考研复试可进半合格，不充合格条数。\n"
                f"{extra}"
                f"业余印象门槛 {BAR} 条合格；本夹合格条数见 Excel。\n",
                encoding="utf-8",
            )
    return by_cell


def cell_stats(by_cell):
    rows = []
    for college in COLLEGES:
        for dept, level_name in (("本科部", "本科部"), ("研究生部", "研究生院")):
            items = by_cell.get((college, dept), [])[:FOLDER_MAX]
            ok = [x for x in items if x["weight"] == "合格"]
            half = [x for x in items if x["weight"] == "半合格"]
            people_ok = {x["people"] for x in ok}
            na = college == "高级翻译学院" and dept == "本科部"
            bar = 0 if na else BAR
            remain = 0 if na else max(0, bar - len(ok))
            remain_p = 0 if na else max(0, PEOPLE_BAR - len(people_ok))
            folder_n = len(items)
            remain_folder = 0 if na else max(0, FOLDER_MIN - folder_n)
            enough = "不适用" if na else ("否" if remain else "是")
            called = "；".join(x["title"] for x in items) or "（无）"
            if na:
                gap = "无普通本科就读对象。余量记 0，不要用保研文件凑数。"
            elif not ok and not half:
                gap = f"合格 0。还差 {remain} 条独立就读自述才够业余印象门槛。距每部 {FOLDER_MIN} 条还差 {remain_folder}。"
            elif not ok:
                gap = f"只有半合格。还差 {remain} 条合格自述（门槛 {bar}）。距每部 {FOLDER_MIN} 条还差 {remain_folder}。"
            else:
                gap = (
                    f"合格 {len(ok)} 条 / {len(people_ok)} 人。还差 {remain} 条、{remain_p} 个独立叙述者。"
                    f"本夹 {folder_n} 条，距每部 {FOLDER_MIN} 条还差 {remain_folder}。"
                )
            rows.append(
                {
                    "学院": college,
                    "层次": level_name,
                    "业余印象门槛_条": bar,
                    "合格自述_条": len(ok),
                    "半合格_条": len(half),
                    "本夹总条数": folder_n,
                    "距7条还差": remain_folder,
                    "独立叙述人数_合格": len(people_ok),
                    "还剩合格条数": remain,
                    "还剩独立叙述者": remain_p,
                    "够业余印象吗": enough,
                    "已调用": called,
                    "缺口": gap,
                }
            )
    return rows


def write_excel(rows, fenyun_split):
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="B8C4CE"),
        right=Side(style="thin", color="B8C4CE"),
        top=Side(style="thin", color="B8C4CE"),
        bottom=Side(style="thin", color="B8C4CE"),
    )
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
    wrap = Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "还剩多少"
    headers = [
        "学院",
        "层次",
        "业余印象门槛（合格条）",
        "已调用合格自述（条）",
        "半合格（条）",
        "本夹总条数",
        "距每部下限7还差",
        "合格独立叙述人数",
        "还剩合格条数",
        "还剩独立叙述者",
        "够一般业余印象分析吗",
        "已调用的自述类文章",
        "缺口说明",
    ]
    ws.append(headers)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
    keys = [
        "学院",
        "层次",
        "业余印象门槛_条",
        "合格自述_条",
        "半合格_条",
        "本夹总条数",
        "距7条还差",
        "独立叙述人数_合格",
        "还剩合格条数",
        "还剩独立叙述者",
        "够业余印象吗",
        "已调用",
        "缺口",
    ]
    for r in rows:
        ws.append([r[k] for k in keys])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=13):
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="宋体", size=10)
            if cell.column in (7, 9, 10) and isinstance(cell.value, int) and cell.value > 0:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
            if cell.column == 11 and cell.value == "否":
                cell.fill = PatternFill("solid", fgColor="F8CBAD")
            if cell.column == 11 and cell.value == "不适用":
                cell.fill = PatternFill("solid", fgColor="D9E2F3")
    widths = [14, 12, 18, 16, 12, 12, 16, 16, 14, 16, 18, 42, 46]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 48
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    ws2 = wb.create_sheet("论坛纷纭再放置")
    ws2.append(["学院", "层次文件夹", "已再放置条数", "距每部7条还差", "文件"])
    for col in range(1, 6):
        c = ws2.cell(1, col)
        c.fill = head_fill
        c.font = head_font
        c.border = thin
    for college, deps in fenyun_split.items():
        for dept, info in deps.items():
            ws2.append([college, f"{college}{dept}", info["placed"], info["short_of_7"], "；".join(info["files"])])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=5):
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="宋体", size=10)
    for i, w in enumerate([14, 22, 14, 16, 70], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("不计入自述权重")
    ws3.append(["学院", "层次", "材料", "为何不算第一人称就读自述"])
    for col in range(1, 5):
        c = ws3.cell(1, col)
        c.fill = head_fill
        c.font = head_font
        c.border = thin
    for row in EXCLUDED_PROMO:
        ws3.append(list(row))
    for r in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=4):
        for cell in r:
            cell.alignment = wrap
            cell.border = thin
            cell.font = Font(name="宋体", size=10)
    for i, w in enumerate([16, 12, 36, 40], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("口径")
    ws4["A1"] = "还剩多少第一人称视角の就读体验"
    ws4["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    notes = [
        f"检索日 {RETRIEVED}（含自述权重+在读续搜）。门槛：每院每层次 {BAR} 条「合格」第一人称就读，且尽量 {PEOPLE_BAR} 个不同叙述者，才够做一般业余印象分析。",
        f"文件夹目标：每部 {FOLDER_MIN}–{FOLDER_MAX} 条第一人称（合格+半合格）。公开取样不够则如实留空，不编造知乎/豆瓣正文。",
        "合格：自称该院该层次在读/刚毕业，写课业、同学、压力、后悔等现场；非约稿、非机构广告。",
        "半合格：身份对得上，但正文是考研/复试/课表转载/入学后两三句；或约稿里密课业第一人称。进夹凑条数，不减合格门槛。",
        "Insight、中青报、张旭/文苑/尚姝辰、赵梦雪、Anette 灯塔：用户已指出可能是宣传出色个例，只进半合格。",
        "就读文章目录：只有第5、6篇 Whever 进入俄语学院本科部。第1篇未写学院，第7篇是咨询号，第2–4篇不是这五院。",
        "高翻本科部不适用。法语学院口译在读帖不记入高翻。出国培训部、iMAP、冠城EAP 不记入统招。",
        "豆瓣外交科普「亚运村小樱花」未确认北外国关，不进夹。",
        "本表在 files/论坛风云/。五个院的部级文件夹同时存在于 论坛纷纭（原取样再放置）和 论坛风云/自述权重+（过滤后）。",
    ]
    for i, line in enumerate(notes, start=3):
        ws4[f"A{i}"] = line
        ws4[f"A{i}"].alignment = wrap
        ws4.row_dimensions[i].height = 32
    ws4.column_dimensions["A"].width = 110

    out = ROOT / "还剩多少第一人称视角の就读体验.xlsx"
    wb.save(out)
    return out


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    WEIGHT.mkdir(parents=True, exist_ok=True)
    fenyun_split = copy_fenyun_split()
    by_cell = write_weight_files()
    rows = cell_stats(by_cell)
    xlsx = write_excel(rows, fenyun_split)
    (ROOT / "检索说明.txt").write_text(
        "论坛风云\n"
        "与「论坛纷纭」并列：纷纭保留原取样不动正文；风云做自述加权和缺口表。\n"
        "自述权重+/学院/学院名本科部|研究生部：合格或半合格第一人称，目标每部 7–10 条，如有才放，不编造。\n"
        "约稿密课业可进半合格以凑条数，不充合格门槛。高翻本科不适用。\n"
        "论坛纷纭/学院/学院名本科部|研究生部：原取样再放置，每部最多 7 条（如有）。\n"
        "Excel：还剩多少第一人称视角の就读体验.xlsx\n"
        "续搜日与检索日同为 2026-09-16。法语学院口译、出国培训部、iMAP 仍不记入五院统招。\n",
        encoding="utf-8",
    )
    (WEIGHT / "口径.txt").write_text(
        "自述权重+ 口径\n"
        "进夹：第一人称、对该院该层次。目标每部 7–10 条，公开取样不够则留空。\n"
        "合格：独立论坛/书评就读现场，非约稿、非机构。\n"
        "半合格：考研/复试/课表转载/入学后两三句；或约稿里密课业第一人称。不充合格条数。\n"
        "不进夹：建制页、他院口译、留学预科、未确认学院的学校层帖、编造正文。\n"
        "高翻本科：不适用。\n",
        encoding="utf-8",
    )
    (ROOT / "目录.json").write_text(
        json.dumps(
            {
                "title": "论坛风云",
                "retrieved": RETRIEVED,
                "excel": "还剩多少第一人称视角の就读体验.xlsx",
                "bar": BAR,
                "folder_min": FOLDER_MIN,
                "folder_max": FOLDER_MAX,
                "cells": rows,
                "fenyun_split": fenyun_split,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(xlsx)
    for r in rows:
        print(
            f"{r['学院']}{r['层次']}: 本夹{r['本夹总条数']} 合格{r['合格自述_条']} "
            f"半{r['半合格_条']} 距7差{r['距7条还差']} 够吗={r['够业余印象吗']}"
        )
    for college, deps in fenyun_split.items():
        for dept, info in deps.items():
            print(f"再放置 {college}{dept}: {info['placed']} 差{info['short_of_7']}")


if __name__ == "__main__":
    main()
