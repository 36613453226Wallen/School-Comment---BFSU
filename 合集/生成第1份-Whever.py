#!/usr/bin/env python3
"""Generate the Whever homepage digest DOCX from the FireShot PDF."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "files" / "口令总表"))
from docx_style import (  # noqa: E402
    INK,
    NAVY,
    SLATE,
    add_bottom_border,
    compact,
    set_run_font,
)

ROOT = Path(__file__).resolve().parent
SOURCE_URL = "https://www.zhihu.com/people/guo-hong-jun-18"
TITLE = "Whever-北外俄语围城后悔录"
ORIGINAL_NAME = f"1{TITLE}.docx"

ANSWERS = [
    {
        "q": "北外上外还值得读吗？语言类大学的头部院校也会失业吗？",
        "date": "发布于 2025-08-23 08:15",
        "folded": False,
        "text": "都会失业，所有人都一样，每年毕业的找不到工作的人大把，我的建议是，除非家里有矿，可以一辈子不用工作，那尽量还是别报小语种。",
    },
    {
        "q": "北外学习压力大吗？",
        "date": "发布于 2025-08-11 08:11",
        "folded": False,
        "text": "太大了",
    },
    {
        "q": "河南农村女要不要报北大小语种提前批（高考全省排名196）？",
        "date": "",
        "folded": True,
        "text": "不要，小语种语感要很好，和高考成绩没关系小语种前途很差，不然会后悔一辈子",
    },
    {
        "q": "我应该报俄语吗？",
        "date": "",
        "folded": True,
        "text": "不要选，不要选，不要选",
    },
    {
        "q": "俄语好学吗 我英语不清楚怎么样 有点迷茫？",
        "date": "发布于 2024-08-06 14:21",
        "folded": False,
        "text": "我英语高考139，一样现在学俄语一窍不通，好好学英语，不会游泳的人换个池子一样淹死",
    },
    {
        "q": "怎样评价巴黎奥运会开幕式？",
        "date": "",
        "folded": True,
        "text": "现在有人出来说我们开始进行八股文考试（提升传统文化，当代人不明白历史等等），开始缠足裹脚（造型美观，走路好看，不喜欢的人都是歧视，要尊重），开始三妻四妾（挽救生育率，别人的自由，你们应该祝福，不应该干涉）。。。。等等不一而足，这就是我觉得巴黎奥运会的内容，讲一万种好处和各种美化，也掩盖不了，这是异化分裂人群，强化阶级差…",
    },
    {
        "q": "不喜欢汉语言文学但家人和机构一直在劝，我该怎么办？",
        "date": "",
        "folded": True,
        "text": "机构太明智了，汉语言文学和法学你要是文科的话就选这两个其他的别想了。听机构的话，因为你根本不知道你自己喜欢什么。听机构的话能少走弯路。你学了才知道你喜欢的其实不喜欢，不然你要费很大劲，转专业代价太大。",
    },
    {
        "q": "大一 法语专业要不要转汉语言文学专业？",
        "date": "",
        "folded": True,
        "text": "家里没钱赶紧转，想考公考央国企赶紧转 不想出国留学硕士转生赶紧转…",
    },
    {
        "q": "求考研俄语语言文学推荐学校？",
        "date": "",
        "folded": True,
        "text": "早点换方向吧，考研俄语只会让你更痛苦，我看我室友考研都退一层皮，早点撤退工作换方向，找找真正喜欢的方向",
    },
    {
        "q": "大学非俄语专业的话，有必要考俄语四级吗？",
        "date": "",
        "folded": True,
        "text": "公共四级随便考着玩，专业四级就算了，突击不出来的",
    },
    {
        "q": "考上北京外国语大学是什么感觉？",
        "date": "",
        "folded": True,
        "text": "不知道大家怎么想到，可能就是围城吧，对我来说，进北外是我人生中最后悔的一件事情，学校不仅对未来毫无保证，而且放弃了大学探索自己的机会，全被无聊繁重的俄语学习占满，并且很清楚，毕业就再也不会再用，现在不过是在混日子罢了。",
    },
    {
        "q": "我是高一纯理科生，分科选择外语的时候选了小语种俄语，对于未来大学选择专业有限制吗？有要求吗？",
        "date": "",
        "folded": True,
        "text": "完蛋了朋友，后面只能去外国语学校，后面毕业即失业，在外国孤身一人天天出差，啃xne6的苦日子等着你呢。…",
    },
    {
        "q": "第一外国语日/德/俄选哪个好呀？",
        "date": "",
        "folded": True,
        "text": "这和你喜欢不喜欢没有关系，直接选日语。选别的等着受折磨",
    },
    {
        "q": "俄语，葡萄牙语，西班牙语，波兰语，这四种语言，哪个比较容易学？",
        "date": "",
        "folded": True,
        "text": "那必是西班牙语最多使用场景，最好学，波兰语，俄语，学几句玩玩行，真认真学痛苦死你",
    },
    {
        "q": "你在COC跑团过程中遇到过最菜的KP能有多菜？",
        "date": "",
        "folded": True,
        "text": "我就是最菜的kp,无论是玩COC，还是玩dnd，可怜到自己只跑过一次coc，后面就从来只有带本，没有玩本的分，想在群里培养几个KP，也从来没有成功过，原本感觉自己描述场景的能力很弱，但是现在学习用AI辅助我创作场景，一下子就支楞起来了…",
    },
    {
        "q": "为了北京外国语大学（211）放弃四川大学（985）值吗？",
        "date": "",
        "folded": True,
        "text": "看到外交部，一看到喜欢英语，一看到北京，地理位置。我就知道味儿对了。且过着吧，有的是你后悔的日子，49年入国军的同学。",
    },
    {
        "q": "汉语言和艺术类选哪个？",
        "date": "",
        "folded": True,
        "text": "汉语言汉语言汉语言，就算喜欢艺术也选汉语言，有份工作比什么都重要，此贴终结",
    },
    {
        "q": "家境一般，录到了北外法语系，应该复读吗？",
        "date": "",
        "folded": True,
        "text": "你认知是没错的，山东大学绝对是更优项，但至少法语学习难度不是那么的大，你还有时间来探索自我，但是变现能力实在是堪忧，外交学就更不提了，外交部，普通人是去不了的，既来之则安之吧，好好学法语，对人的气质有很大的提升。",
    },
    {
        "q": "我马上高二想学俄语 等到我俄语高考人会变很多吗？",
        "date": "",
        "folded": True,
        "text": "想好就行，后面学乌克兰语哈萨克斯坦语别哭就好",
    },
    {
        "q": "大学学的俄语专业，该不该放弃？",
        "date": "",
        "folded": True,
        "text": "早日放弃，选择比努力重要的多，方向不对，再辛苦也没用",
    },
    {
        "q": "英语不好，推荐学俄语代替英语高考吗？",
        "date": "",
        "folded": True,
        "text": "见仁见智吧。反正你如果高考考完上大学这个专业，相当于是高起点。但实话说，俄语，前途非常窄。英语确实是世界语言，所以尽可能要学英语。当然如果你是从功利的角度俄语出题出的比较简单，但俄语学习难度也是很高的。而且。我知道的，山西那边的俄语，都是从初中开始就有学俄语的，你高中再开始学有一点迟了",
    },
]


def new_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    return doc


def add_kicker(doc, text):
    p = doc.add_paragraph()
    compact(p, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, "黑体", 10, True, NAVY)


def add_title(doc, text):
    p = doc.add_paragraph()
    compact(p, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_bottom_border(p)
    run = p.add_run(text)
    set_run_font(run, "黑体", 16, True, NAVY)


def add_meta(doc):
    lines = [
        "作者：Whever　签名：京城苦难自闭小语种大学生",
        "知乎主页：https://www.zhihu.com/people/guo-hong-jun-18",
        "首次发布：截图可见最早一条为 2024-08-06 14:21；其余多数未显示，不补造",
        "编辑时间：未知（主页列表未给出）",
        "PDF 快照：2026-09-15 11:16 UTC（FireShot 元数据）",
        "学校 / 专业：北京外国语大学 · 俄语（小语种，据回答自述）",
        "说明：这是主页回答列表的合订摘录，不是单篇全文。标了「阅读全文」的只保留截图可见部分。已去掉导航、私信数、赞同条和评论区。",
    ]
    for line in lines:
        p = doc.add_paragraph()
        compact(p, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(line)
        set_run_font(run, "宋体", 9, color=SLATE)


def add_answers(doc):
    for i, item in enumerate(ANSWERS, start=1):
        h = doc.add_paragraph()
        compact(h, before=12, after=4, line=1.2)
        add_bottom_border(h, "8A4B12", "8")
        run = h.add_run(f"{i}. {item['q']}")
        set_run_font(run, "黑体", 12, True, NAVY)

        if item["date"]:
            d = doc.add_paragraph()
            compact(d, after=3, line=1.1)
            run = d.add_run(item["date"])
            set_run_font(run, "宋体", 9, color=SLATE)

        if item["folded"]:
            tag = doc.add_paragraph()
            compact(tag, after=3, line=1.1)
            run = tag.add_run("截图未展开全文，下面只是列表摘要。")
            set_run_font(run, "楷体", 9, color=SLATE)

        p = doc.add_paragraph()
        compact(p, after=6, line=1.35, first_line=22)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        run = p.add_run(item["text"])
        set_run_font(run, "宋体", 12, color=INK)


def add_link(doc):
    p = doc.add_paragraph()
    compact(p, before=12, after=0, line=1.15)
    add_bottom_border(p, "D6DEE8", "8")
    run = p.add_run("知乎主页：")
    set_run_font(run, "黑体", 10, True, NAVY)
    run = p.add_run(SOURCE_URL)
    set_run_font(run, "宋体", 10, color=RGBColor(0x2E, 0x74, 0xB5))


def write_original():
    doc = new_doc()
    add_kicker(doc, "合集  ·  作者主页摘录")
    add_title(doc, TITLE)
    add_meta(doc)
    add_answers(doc)
    add_link(doc)
    path = ROOT / ORIGINAL_NAME
    doc.core_properties.title = TITLE
    doc.save(path)
    print(f"已生成：{path}")
    return path


if __name__ == "__main__":
    write_original()
